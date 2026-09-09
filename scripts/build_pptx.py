#!/usr/bin/env python3
"""Génère la présentation exécutive CMDT AI au format PowerPoint 16:9.

Prérequis :  pip install python-pptx
Usage     :  python scripts/build_pptx.py [sortie.pptx]
             (défaut : cmdt-ai-presentation.pptx à la racine du projet)
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

sys.path.insert(0, str(Path(__file__).resolve().parent))
from content import SLIDES, PRODUCT, CHAIN, NAVY, LEAF, LEAF_SOFT, INK, INK_MUTED, WHITE, CANVAS  # noqa: E402

EMU_W, EMU_H = Inches(13.333), Inches(7.5)


def rgb(t):
    return RGBColor(*t)


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()


def _box(slide, x, y, w, h):
    from pptx.enum.shapes import MSO_SHAPE

    return slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)


def _text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, size, bold, color) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = "Inter"
        r.font.color.rgb = rgb(color)
    return tb


def _logo(slide, x, y, on_dark=False):
    from pptx.enum.shapes import MSO_SHAPE

    sq = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.42), Inches(0.42))
    _fill(sq, LEAF)
    sq.text_frame.text = "C"
    r = sq.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = rgb(NAVY)
    r.font.name = "Georgia"
    sq.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    _text(
        slide,
        x + Inches(0.55),
        y - Inches(0.04),
        Inches(4),
        Inches(0.5),
        [("CMDT AI", 13, True, WHITE if on_dark else NAVY)],
    )


def _footer(slide, txt):
    _text(
        slide,
        Inches(0.9),
        Inches(7.0),
        Inches(11.5),
        Inches(0.4),
        [(txt, 9, False, INK_MUTED)],
    )


def slide_cover(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = _box(slide, 0, 0, EMU_W, EMU_H)
    _fill(bg, NAVY)
    _logo(slide, Inches(0.9), Inches(0.7), on_dark=True)
    _text(
        slide,
        Inches(0.9),
        Inches(2.6),
        Inches(11.5),
        Inches(2.2),
        [
            (s["title"], 54, True, WHITE),
            (s["subtitle"], 24, False, LEAF_SOFT),
        ],
    )
    _text(slide, Inches(0.9), Inches(5.3), Inches(11), Inches(0.5), [(s["note"], 14, False, WHITE)])
    _text(
        slide,
        Inches(0.9),
        Inches(6.9),
        Inches(11.5),
        Inches(0.4),
        [(f"{PRODUCT['org']}  ·  {PRODUCT['version']}  ·  {PRODUCT['disclaimer']}", 9, False, LEAF_SOFT)],
    )


def _header(slide, kicker, title):
    _logo(slide, Inches(0.9), Inches(0.55))
    _text(slide, Inches(0.9), Inches(1.25), Inches(11.5), Inches(0.4),
          [(kicker.upper(), 11, True, LEAF)])
    _text(slide, Inches(0.9), Inches(1.6), Inches(11.5), Inches(1.0),
          [(title, 30, True, NAVY)])
    bar = _box(slide, Inches(0.9), Inches(2.55), Inches(1.4), Pt(3))
    _fill(bar, LEAF)


def slide_content(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill(_box(slide, 0, 0, EMU_W, EMU_H), CANVAS)
    _header(slide, s["kicker"], s["title"])
    runs = []
    for b in s["bullets"]:
        runs.append(("•  " + b, 15, False, INK))
    _text(slide, Inches(0.9), Inches(3.0), Inches(11.5), Inches(3.6), runs, spacing=1.25)
    if s.get("footer"):
        _footer(slide, s["footer"])


def slide_chain(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill(_box(slide, 0, 0, EMU_W, EMU_H), CANVAS)
    _header(slide, s["kicker"], s["title"])
    runs = [("•  " + b, 15, False, INK) for b in s["body"]]
    _text(slide, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.8), runs, spacing=1.25)
    x = Inches(0.9)
    y = Inches(5.2)
    for i, step in enumerate(CHAIN):
        w = Inches(1.55)
        chip = _box(slide, x, y, w, Inches(0.55))
        _fill(chip, NAVY if i % 2 == 0 else LEAF)
        chip.text_frame.text = step
        pr = chip.text_frame.paragraphs[0]
        pr.alignment = PP_ALIGN.CENTER
        pr.runs[0].font.size = Pt(11)
        pr.runs[0].font.bold = True
        pr.runs[0].font.color.rgb = rgb(WHITE)
        x += w + Inches(0.12)
    _footer(slide, PRODUCT["url"])


def slide_module(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill(_box(slide, 0, 0, EMU_W, EMU_H), CANVAS)
    side = _box(slide, 0, 0, Inches(2.4), EMU_H)
    _fill(side, NAVY)
    _text(slide, Inches(0.35), Inches(0.8), Inches(1.8), Inches(1.0),
          [(s["no"], 44, True, LEAF_SOFT)])
    _text(slide, Inches(0.35), Inches(2.0), Inches(1.9), Inches(3.0),
          [("MODULE", 10, True, WHITE)])
    _text(slide, Inches(2.9), Inches(0.7), Inches(9.8), Inches(1.0),
          [(s["name"], 26, True, NAVY)])
    blocks = [
        ("CONSTAT", s["what"]),
        ("POURQUOI", s["why"]),
        ("ENJEU", s["sowhat"]),
        ("ACTION", s["nowwhat"]),
    ]
    y = Inches(1.9)
    for label, body in blocks:
        _text(slide, Inches(2.9), y, Inches(9.8), Inches(0.3), [(label, 11, True, LEAF)])
        _text(slide, Inches(2.9), y + Inches(0.32), Inches(9.8), Inches(1.0),
              [(body, 13, False, INK)], spacing=1.2)
        y += Inches(1.25)
    _footer(slide, f"{PRODUCT['name']} · {PRODUCT['tagline']}")


BUILDERS = {
    "cover": slide_cover,
    "content": slide_content,
    "chain": slide_chain,
    "module": slide_module,
}


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "cmdt-ai-presentation.pptx"
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    for s in SLIDES:
        BUILDERS[s["type"]](prs, s)
    prs.save(str(out))
    print(f"OK  {out}  ({len(SLIDES)} diapositives)")


if __name__ == "__main__":
    main()
